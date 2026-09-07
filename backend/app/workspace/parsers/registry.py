# backend/app/workspace/parsers/registry.py
import logging
import csv
import io
from pathlib import Path
from typing import Optional
from docx import Document
import openpyxl
import pypdfium2 as pdfium

logger = logging.getLogger(__name__)

# Step 8.2: Lazy-Load the OCR engine to prevent massive VRAM/RAM spikes on startup
_ocr_engine_instance = None
_ocr_engine_loaded = False

def get_ocr_engine():
    global _ocr_engine_instance, _ocr_engine_loaded
    if not _ocr_engine_loaded:
        _ocr_engine_loaded = True
        try:
            from paddleocr import PaddleOCR
            # FIX: Removed invalid 'show_log=False' argument. Initializes only when requested.
            _ocr_engine_instance = PaddleOCR(use_angle_cls=True, lang='en')
        except ImportError:
            logger.warning("PaddleOCR not available. Scanned PDFs will fail extraction.")
    return _ocr_engine_instance

class DocumentParserRegistry:
    @staticmethod
    def parse(file_path: Path, mime_type: str) -> str:
        """Routes a file to the appropriate native parser or OCR engine."""
        
        if mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return DocumentParserRegistry._parse_docx(file_path)
            
        elif mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            return DocumentParserRegistry._parse_xlsx(file_path)
            
        elif mime_type == "text/csv" or mime_type == "text/plain":
            return DocumentParserRegistry._parse_csv(file_path)
            
        elif mime_type == "application/pdf":
            return DocumentParserRegistry._parse_pdf(file_path)
            
        elif mime_type.startswith("image/"):
            return DocumentParserRegistry._parse_image(file_path)
            
        else:
            raise ValueError(f"No parser available for mime type: {mime_type}")

    @staticmethod
    def _parse_docx(file_path: Path) -> str:
        doc = Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs if para.text.strip()])

    @staticmethod
    def _parse_xlsx(file_path: Path) -> str:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"--- Sheet: {sheet.title} ---")
            for row in sheet.iter_rows(values_only=True):
                row_str = " | ".join([str(cell) for cell in row if cell is not None])
                if row_str.strip():
                    lines.append(row_str)
        return "\n".join(lines)

    @staticmethod
    def _parse_csv(file_path: Path) -> str:
        lines = []
        with file_path.open("r", encoding="utf-8") as f:
            reader = csv.reader(f)
            for row in reader:
                lines.append(" | ".join(row))
        return "\n".join(lines)

    @staticmethod
    def _parse_pdf(file_path: Path) -> str:
        pdf = pdfium.PdfDocument(str(file_path))
        extracted_pages = []
        
        for i in range(len(pdf)):
            page = pdf[i]
            text_page = page.get_textpage()
            native_text = text_page.get_text_range()
            
            # Step 8.2: PaddleOCR Fallback for scanned pages
            if not native_text or len(native_text.strip()) < 20:
                engine = get_ocr_engine()
                if engine:
                    logger.info(f"Page {i} lacks native text. Falling back to PaddleOCR.")
                    bitmap = page.render(scale=2)
                    pil_image = bitmap.to_pil()
                    
                    img_byte_arr = io.BytesIO()
                    pil_image.save(img_byte_arr, format='JPEG')
                    img_bytes = img_byte_arr.getvalue()
                    
                    result = engine.ocr(img_bytes, cls=True)
                    ocr_text = []
                    if result and result[0]:
                        for line in result[0]:
                            ocr_text.append(line[1][0])
                    extracted_pages.append("\n".join(ocr_text))
                else:
                    extracted_pages.append(f"[Scanned Image - OCR Unavailable]")
            else:
                extracted_pages.append(native_text)
                
        return "\n\n".join(extracted_pages)

    @staticmethod
    def _parse_image(file_path: Path) -> str:
        engine = get_ocr_engine()
        if not engine:
            raise RuntimeError("PaddleOCR is required to extract text from images.")
        
        with file_path.open("rb") as f:
            result = engine.ocr(f.read(), cls=True)
            ocr_text = []
            if result and result[0]:
                for line in result[0]:
                    ocr_text.append(line[1][0])
            return "\n".join(ocr_text)